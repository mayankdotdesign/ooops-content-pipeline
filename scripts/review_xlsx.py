"""
Builds and reads the review spreadsheets for Phase 6's two-cycle
approval flow. Shared by both cycles — the column set differs, the
Status-dropdown mechanics don't.

Usage: imported by scripts/review_cycle.py, not run directly.
"""

import os

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

STATUS_OPTIONS = ["Approved", "Need Change", "Rejected"]


def _slide_summary(item):
    """One-cell text summary of an item's slides, for the content-review
    sheet — a human reading the spreadsheet needs to see what's in each
    slide without opening the (not-yet-rendered) image."""
    parts = []
    for i, slide in enumerate(item["slides"], start=1):
        layout = slide["layout"]
        if layout == "hook":
            parts.append(f"[{i}:hook] {slide['text']}")
        elif layout == "bullet_list":
            parts.append(f"[{i}:bullet_list] " + " / ".join(slide["items"]))
        elif layout == "stat_card":
            parts.append(f"[{i}:stat_card] {slide['stat']} — {slide['caption']}")
        elif layout == "photo":
            parts.append(f"[{i}:photo] {slide['caption']}")
        elif layout == "logo_endcard":
            parts.append(f"[{i}:logo_endcard]")
    return "\n".join(parts)


def build_content_review_sheet(items, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Content Review"
    headers = ["Post ID", "Content/slide text", "Caption", "Hashtags", "CTA", "Post type", "Status", "My Comments"]
    ws.append(headers)

    for item in items:
        ws.append([
            item["id"],
            _slide_summary(item),
            item["caption"],
            " ".join(item["hashtags"]),
            item["cta"],
            item["post_type"],
            "",
            "",
        ])

    _apply_status_dropdown(ws, status_col=7, n_rows=len(items))
    _autosize(ws, widths={1: 8, 2: 60, 3: 40, 4: 30, 5: 40, 6: 12, 7: 14, 8: 40})
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb.save(out_path)
    return out_path


def build_visual_review_sheet(items, image_urls_by_id, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Visual Review"
    headers = ["Post ID", "Image link(s)", "Caption (context)", "Status", "My Comments"]
    ws.append(headers)

    for item in items:
        urls = image_urls_by_id[item["id"]]
        ws.append([
            item["id"],
            "\n".join(urls),
            item["caption"],
            "",
            "",
        ])

    _apply_status_dropdown(ws, status_col=4, n_rows=len(items))
    _autosize(ws, widths={1: 8, 2: 60, 3: 40, 4: 14, 5: 40})
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb.save(out_path)
    return out_path


def _apply_status_dropdown(ws, status_col, n_rows):
    dv = DataValidation(type="list", formula1=f'"{",".join(STATUS_OPTIONS)}"', allow_blank=True)
    ws.add_data_validation(dv)
    col_letter = get_column_letter(status_col)
    dv.add(f"{col_letter}2:{col_letter}{n_rows + 1}")


def _autosize(ws, widths):
    for col_idx, width in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = cell.alignment.copy(wrap_text=True, vertical="top")


def read_review_sheet(path):
    """Returns a list of dicts: {post_id, status, comments} for every row
    that has a Post ID, regardless of sheet layout (content vs visual —
    both put Post ID first and Status/My Comments as the last two
    non-blank-required columns, found by header name so column order
    doesn't matter if the reviewer rearranges anything)."""
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    header_row = [cell.value for cell in ws[1]]
    col_index = {name: i for i, name in enumerate(header_row) if name}

    results = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[col_index["Post ID"]] is None:
            continue
        results.append({
            "post_id": int(row[col_index["Post ID"]]),
            "status": (row[col_index["Status"]] or "").strip(),
            "comments": (row[col_index["My Comments"]] or "").strip(),
        })
    return results
